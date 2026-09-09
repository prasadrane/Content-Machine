"""IdeaBankConnector: Ingestion and export connector for Prasad Windmill Idea Bank Excel."""

from __future__ import annotations

import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

import openpyxl

from .base import BaseConnector, ConnectorItem, ConnectorResult

logger = logging.getLogger(__name__)


class IdeaBankConnector(BaseConnector):
    """Two-way connector for reading from and writing to Windmill Idea Bank Excel workbook."""

    def __init__(self, file_path: str | Path | None = None):
        self.file_path = Path(file_path) if file_path is not None else None

    def _resolve_path(self, override_path: str | Path | None = None) -> Path:
        target = Path(override_path) if override_path is not None else self.file_path
        if target is None:
            raise ValueError("No Excel file path provided to IdeaBankConnector.")
        if not target.exists():
            raise FileNotFoundError(f"Idea Bank Excel file not found at: {target}")
        return target

    def fetch(self, file_path: str | Path | None = None, **kwargs) -> ConnectorResult:
        """Read ideas from the 'IDEAS' sheet and return ConnectorItem list."""
        path = self._resolve_path(file_path)
        wb = openpyxl.load_workbook(path, data_only=True)

        if "IDEAS" not in wb.sheetnames:
            raise ValueError(f"Sheet 'IDEAS' not found in workbook: {path}")

        ws = wb["IDEAS"]
        items: list[ConnectorItem] = []
        now = datetime.now(timezone.utc)

        for row in range(2, ws.max_row + 1):
            val_title = ws.cell(row=row, column=1).value
            val_brainstorm = ws.cell(row=row, column=2).value

            title = str(val_title or "").strip()
            body = str(val_brainstorm or "").strip()

            if not title and not body:
                continue

            item_title = title if title else body[:60]
            item_body = body if body else title

            items.append(
                ConnectorItem(
                    title=item_title,
                    url=f"file://{path.resolve()}#row={row}",
                    body=item_body,
                    source="idea_bank",
                    fetched_at=now,
                    guid=f"idea_bank:{row}:{item_title[:30]}",
                )
            )

        return ConnectorResult(source="idea_bank", items=items)

    def append_idea(
        self,
        title: str,
        brainstorm: str = "",
        file_path: str | Path | None = None,
    ) -> int:
        """Append a new idea row to the 'IDEAS' sheet and save."""
        path = self._resolve_path(file_path)
        wb = openpyxl.load_workbook(path)
        if "IDEAS" not in wb.sheetnames:
            ws = wb.create_sheet(title="IDEAS")
            ws.append(["Idea Title", "Idea Brainstorm", None, "Check", "Idea Checklist:"])
        else:
            ws = wb["IDEAS"]

        # Find the next available row where column 1 and 2 are empty (starting from row 2)
        next_row = 2
        for r in range(2, ws.max_row + 2):
            if not ws.cell(row=r, column=1).value and not ws.cell(row=r, column=2).value:
                next_row = r
                break

        ws.cell(row=next_row, column=1, value=title.strip())
        ws.cell(row=next_row, column=2, value=brainstorm.strip())

        wb.save(path)
        return next_row

    def append_packaging(
        self,
        final_title: str,
        brainstorm: str = "",
        template_inspiration: str = "",
        thumbnail_rec: str = "",
        file_path: str | Path | None = None,
    ) -> int:
        """Append a new packaging entry to the 'PACKAGING' sheet and save."""
        path = self._resolve_path(file_path)
        wb = openpyxl.load_workbook(path)
        if "PACKAGING" not in wb.sheetnames:
            ws = wb.create_sheet(title="PACKAGING")
            ws.append(["Final Title", "Idea Brainstorm", None, "Title Template Inspiration", "Thumbnail Recommendation"])
        else:
            ws = wb["PACKAGING"]

        next_row = 2
        for r in range(2, ws.max_row + 2):
            if not ws.cell(row=r, column=1).value and not ws.cell(row=r, column=2).value:
                next_row = r
                break

        ws.cell(row=next_row, column=1, value=final_title.strip())
        ws.cell(row=next_row, column=2, value=brainstorm.strip())
        if template_inspiration:
            ws.cell(row=next_row, column=4, value=template_inspiration.strip())
        if thumbnail_rec:
            ws.cell(row=next_row, column=5, value=thumbnail_rec.strip())

        wb.save(path)
        return next_row

    def export_idea_batch(
        self,
        ideas: list[dict[str, str]],
        file_path: str | Path | None = None,
    ) -> int:
        """Export multiple ideas in a single open-and-save operation."""
        path = self._resolve_path(file_path)
        wb = openpyxl.load_workbook(path)
        ws = wb["IDEAS"] if "IDEAS" in wb.sheetnames else wb.active

        # Find first empty row
        next_row = 2
        for r in range(2, ws.max_row + 2):
            if not ws.cell(row=r, column=1).value and not ws.cell(row=r, column=2).value:
                next_row = r
                break

        count = 0
        for item in ideas:
            t = item.get("title", "").strip()
            b = item.get("brainstorm", "").strip()
            if t or b:
                ws.cell(row=next_row, column=1, value=t)
                ws.cell(row=next_row, column=2, value=b)
                next_row += 1
                count += 1

        wb.save(path)
        return count

    def export_packaging_batch(
        self,
        packagings: list[dict[str, str]],
        file_path: str | Path | None = None,
    ) -> int:
        """Export multiple packaging entries in a single open-and-save operation."""
        path = self._resolve_path(file_path)
        wb = openpyxl.load_workbook(path)
        ws = wb["PACKAGING"] if "PACKAGING" in wb.sheetnames else wb.create_sheet(title="PACKAGING")

        next_row = 2
        for r in range(2, ws.max_row + 2):
            if not ws.cell(row=r, column=1).value and not ws.cell(row=r, column=2).value:
                next_row = r
                break

        count = 0
        for item in packagings:
            ft = item.get("final_title", "").strip()
            b = item.get("brainstorm", "").strip()
            ti = item.get("template_inspiration", "").strip()
            tr = item.get("thumbnail_rec", "").strip()
            if ft or b:
                ws.cell(row=next_row, column=1, value=ft)
                ws.cell(row=next_row, column=2, value=b)
                if ti:
                    ws.cell(row=next_row, column=4, value=ti)
                if tr:
                    ws.cell(row=next_row, column=5, value=tr)
                next_row += 1
                count += 1

        wb.save(path)
        return count

