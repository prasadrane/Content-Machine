"""Unit tests for the imaging pipeline: per-post §28 image-prompt generation.

Covers: post loading from the Windmill xlsx (IDEAS + PACKAGING alignment),
LLM template generation with validation/retry, fabricated-number flagging,
parallel batch generation, PACKAGING column F write-back with backup, and
consolidated markdown export.

Follows strict TDD principles. Tests run offline (fake router).
"""

from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from content_machine.connectors.idea_bank import IdeaBankConnector
from content_machine.imaging.prompt_gen import (
    REQUIRED_HEADERS,
    GeneratedImagePrompt,
    ImagePromptGenerator,
    flag_fabricated_numbers,
    load_posts,
    export_markdown,
)


VALID_TEMPLATE = """## POST TITLE
Why One .Result Call Freezes Everything

## POST CONTENT / IDEA
Draft about thread pool starvation; agents hit 429s and partial responses.

## CATEGORY
PERFORMANCE

## CORE ENGINEERING PROBLEM
Thread pool starvation from sync-over-async calls.

## TRIGGER
A single blocking .Result call under load.

## MECHANISM TO VISUALIZE
Blocked threads occupy pool slots; queued work waits.

## CONSEQUENCE
Requests queue, latency spikes, 429 throttling.

## BETTER STATE / SOLUTION
Async all the way down.

## IMPORTANT TECHNICAL LABELS
.Result, ThreadPool, 429

## OPTIONAL TECHNICAL ARTIFACT
(none)

## EXACT FACTS THAT MAY APPEAR
429
"""

DRAFT_TEXT = (
    "How one .Result call starves the thread pool.\n"
    "Agents hit 429s and partial responses; 200 lines of code hide the trap.\n"
    "Async all the way down fixes it."
)


class FakeRouter:
    """Stands in for ModelRouter; records calls, returns canned responses."""

    def __init__(self, responses):
        # responses: list of str/Exception consumed in order, or a callable
        self.responses = responses
        self.calls: list[dict] = []

    def complete(self, model_layer, prompt, *, system=None, schema=None, **kwargs):
        self.calls.append(
            {"model_layer": model_layer, "prompt": prompt, "system": system,
             "schema": schema, "kwargs": kwargs}
        )
        item = self.responses(len(self.calls) - 1) if callable(self.responses) \
            else self.responses[min(len(self.calls) - 1, len(self.responses) - 1)]
        if isinstance(item, Exception):
            raise item
        return item


def make_workbook(path: Path, n_posts: int = 3):
    """Fixture workbook mirroring the real one: IDEAS rows 2..n+1, PACKAGING aligned."""
    wb = openpyxl.Workbook()
    ws_ideas = wb.active
    ws_ideas.title = "IDEAS"
    ws_ideas.append(["Idea Title", "Idea Brainstorm", None, "Check", "Idea Checklist:"])
    ws_pack = wb.create_sheet(title="PACKAGING")
    ws_pack.append(["Final Title", "Idea Brainstorm", None,
                    "Title Template Inspiration", "Thumbnail Recommendation"])

    for i in range(n_posts):
        ws_ideas.append([f"Idea Title {i + 1}", DRAFT_TEXT])
        ws_pack.append([f"Final Title {i + 1}", f"short desc {i + 1}", None,
                        "From X to Y", f"Thumbnail rec {i + 1}"])
    wb.save(path)


class TestLoadPosts(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.xlsx_path = Path(self.temp_dir.name) / "Test_Idea_Bank.xlsx"
        make_workbook(self.xlsx_path, n_posts=3)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_load_posts_reads_aligned_rows(self):
        posts = load_posts(self.xlsx_path)
        self.assertEqual(len(posts), 3)
        p = posts[0]
        self.assertEqual(p.row, 2)
        self.assertEqual(p.idea_title, "Idea Title 1")
        self.assertIn("429s", p.draft)
        self.assertEqual(p.final_title, "Final Title 1")
        self.assertEqual(p.thumbnail_rec, "Thumbnail rec 1")
        self.assertEqual(posts[2].row, 4)

    def test_load_posts_skips_empty_rows(self):
        wb = openpyxl.load_workbook(self.xlsx_path)
        wb["IDEAS"].append([None, None])
        wb["IDEAS"].append(["Valid After Gap", DRAFT_TEXT])
        wb["PACKAGING"].append([None, None])
        wb["PACKAGING"].append(["Valid After Gap Final", "d", None, "t", "r"])
        wb.save(self.xlsx_path)

        posts = load_posts(self.xlsx_path)
        self.assertEqual(len(posts), 4)
        self.assertEqual(posts[3].row, 6)
        self.assertEqual(posts[3].idea_title, "Valid After Gap")

    def test_load_posts_respects_row_bounds(self):
        posts = load_posts(self.xlsx_path, first_row=3, last_row=3)
        self.assertEqual(len(posts), 1)
        self.assertEqual(posts[0].row, 3)


class TestGenerateOne(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.xlsx_path = Path(self.temp_dir.name) / "Test_Idea_Bank.xlsx"
        make_workbook(self.xlsx_path, n_posts=1)
        self.post = load_posts(self.xlsx_path)[0]

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_generate_one_returns_filled_template(self):
        router = FakeRouter([VALID_TEMPLATE])
        gen = ImagePromptGenerator(router=router, model="qwen3.8-max")
        result = gen.generate_one(self.post)

        self.assertIsInstance(result, GeneratedImagePrompt)
        self.assertEqual(result.row, 2)
        for header in REQUIRED_HEADERS:
            self.assertIn(header, result.text)
        self.assertEqual(result.warnings, [])

    def test_generate_one_passes_draft_and_thinking_kwargs(self):
        router = FakeRouter([VALID_TEMPLATE])
        gen = ImagePromptGenerator(router=router, model="qwen3.8-max")
        gen.generate_one(self.post)

        call = router.calls[0]
        self.assertEqual(call["model_layer"], ["qwen3.8-max"])
        self.assertIn("429s", call["prompt"])          # draft forwarded
        self.assertIn("Thumbnail rec 1", call["prompt"])  # thumbnail hint forwarded
        self.assertIn("## POST TITLE", call["system"])  # §28 template in system
        self.assertEqual(call["kwargs"]["thinking"]["type"], "enabled")
        self.assertEqual(call["kwargs"]["extra_body"], {"reasoning_effort": "high"})

    def test_generate_one_retries_once_on_invalid_output(self):
        invalid = VALID_TEMPLATE.replace("## TRIGGER\nA single blocking .Result call under load.\n", "")
        router = FakeRouter([invalid, VALID_TEMPLATE])
        gen = ImagePromptGenerator(router=router, model="qwen3.8-max")
        result = gen.generate_one(self.post)

        self.assertEqual(len(router.calls), 2)
        for header in REQUIRED_HEADERS:
            self.assertIn(header, result.text)

    def test_generate_one_raises_after_second_invalid(self):
        invalid = "no headers at all"
        router = FakeRouter([invalid, invalid])
        gen = ImagePromptGenerator(router=router, model="qwen3.8-max")
        with self.assertRaises(ValueError):
            gen.generate_one(self.post)
        self.assertEqual(len(router.calls), 2)

    def test_generate_one_flags_fabricated_numbers_as_warnings(self):
        # Template invents "1m 48s" — neither 1 nor 48 appear in the draft.
        fabricated = VALID_TEMPLATE.replace(
            "## CONSEQUENCE\nRequests queue, latency spikes, 429 throttling.",
            "## CONSEQUENCE\nRequests queue for 1m 48s, latency spikes, 429 throttling.",
        )
        router = FakeRouter([fabricated])
        gen = ImagePromptGenerator(router=router, model="qwen3.8-max")
        result = gen.generate_one(self.post)

        self.assertTrue(any("48" in w for w in result.warnings))


class TestFlagFabricatedNumbers(unittest.TestCase):

    def test_flags_numbers_absent_from_source(self):
        source = "Agents hit 429s; 200 lines hide the trap."
        generated = "Throttled at 429 after 1m 48s; 200 lines."
        warnings = flag_fabricated_numbers(generated, source)
        flagged = " ".join(warnings)
        self.assertIn("48", flagged)
        self.assertNotIn("429 ", flagged + " ")
        self.assertNotIn("200", flagged)

    def test_no_warnings_when_all_numbers_sourced(self):
        source = "429s and 200 lines, 16-18 threads."
        generated = "429 throttle; 200 lines; 16-18 threads."
        self.assertEqual(flag_fabricated_numbers(generated, source), [])

    def test_ignores_template_structural_numbers(self):
        source = "no numbers here"
        generated = "## POST TITLE\nHeadline\n\n## CATEGORY\nDATABASES\n"
        self.assertEqual(flag_fabricated_numbers(generated, source), [])


class TestGenerateAll(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.xlsx_path = Path(self.temp_dir.name) / "Test_Idea_Bank.xlsx"
        make_workbook(self.xlsx_path, n_posts=3)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_generate_all_returns_result_per_row(self):
        router = FakeRouter([VALID_TEMPLATE])
        gen = ImagePromptGenerator(router=router, model="qwen3.8-max", max_workers=2)
        posts = load_posts(self.xlsx_path)
        results = gen.generate_all(posts)

        self.assertEqual(len(results), 3)
        self.assertEqual([r.row for r in results], [2, 3, 4])
        self.assertEqual(len(router.calls), 3)
        for r in results:
            self.assertIn("## POST TITLE", r.text)


class TestWriteBack(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.xlsx_path = Path(self.temp_dir.name) / "Test_Idea_Bank.xlsx"
        make_workbook(self.xlsx_path, n_posts=3)
        self.connector = IdeaBankConnector(file_path=self.xlsx_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_write_image_prompts_fills_column_f(self):
        prompts = {2: "PROMPT ROW 2", 3: "PROMPT ROW 3"}
        self.connector.write_image_prompts(prompts)

        wb = openpyxl.load_workbook(self.xlsx_path)
        ws = wb["PACKAGING"]
        self.assertEqual(ws.cell(row=1, column=6).value, "Image Prompt")
        self.assertEqual(ws.cell(row=2, column=6).value, "PROMPT ROW 2")
        self.assertEqual(ws.cell(row=3, column=6).value, "PROMPT ROW 3")
        self.assertIsNone(ws.cell(row=4, column=6).value)

    def test_write_image_prompts_creates_backup(self):
        backup_dir = Path(self.temp_dir.name) / "backups"
        backup_dir.mkdir()
        prompts = {2: "PROMPT ROW 2"}
        backup_path = self.connector.write_image_prompts(prompts, backup_dir=backup_dir)

        self.assertIsNotNone(backup_path)
        self.assertTrue(backup_path.exists())
        self.assertIn("pre-imageprompts", backup_path.name)
        # Backup holds pre-write state: no column F
        wb = openpyxl.load_workbook(backup_path)
        self.assertIsNone(wb["PACKAGING"].cell(row=1, column=6).value)

    def test_write_image_prompts_no_backup_when_disabled(self):
        backup_path = self.connector.write_image_prompts({2: "X"}, backup=False)
        self.assertIsNone(backup_path)


class TestExportMarkdown(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.out_path = Path(self.temp_dir.name) / "all_image_prompts.md"

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_export_contains_every_post_fenced(self):
        results = [
            GeneratedImagePrompt(row=2, final_title="Final Title 1",
                                 text=VALID_TEMPLATE, warnings=[]),
            GeneratedImagePrompt(row=3, final_title="Final Title 2",
                                 text=VALID_TEMPLATE, warnings=["invented 48"]),
        ]
        export_markdown(results, self.out_path)

        content = self.out_path.read_text(encoding="utf-8")
        self.assertIn("Final Title 1", content)
        self.assertIn("Final Title 2", content)
        self.assertIn("Row 2", content)
        self.assertIn("```", content)  # fenced for copy-paste fidelity
        self.assertIn("invented 48", content)  # warnings surfaced


if __name__ == "__main__":
    unittest.main(verbosity=2)
