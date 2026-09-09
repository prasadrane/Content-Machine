"""Unit tests for IdeaBankConnector (Windmill Idea Bank Excel integration).

Follows strict TDD principles. Tests run offline.
"""

from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from content_machine.connectors.idea_bank import IdeaBankConnector
from content_machine.connectors.base import ConnectorItem, ConnectorResult


class TestIdeaBankConnector(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.xlsx_path = Path(self.temp_dir.name) / "Test_Idea_Bank.xlsx"
        self._create_mock_workbook(self.xlsx_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def _create_mock_workbook(self, path: Path):
        wb = openpyxl.Workbook()
        # Default sheet -> IDEAS
        ws_ideas = wb.active
        ws_ideas.title = "IDEAS"
        ws_ideas.append(["Idea Title", "Idea Brainstorm", None, "Check", "Idea Checklist:"])
        ws_ideas.append(["Antigravity vs Claude Code Latency", "Comparing tool call speed and developer cognitive flow."])
        ws_ideas.append(["Kafka Consumer Lag War Story", "Debugging lag spikes during sudden partition rebalancing."])

        # PACKAGING sheet
        ws_pack = wb.create_sheet(title="PACKAGING")
        ws_pack.append(["Final Title", "Idea Brainstorm", None, "Title Template Inspiration", "Thumbnail Recommendation"])
        ws_pack.append(["Why Fast Agents Keep You in the Zone", "Antigravity vs Claude Code", None, "From X to Y", "Before/After"])

        # DONE sheet
        ws_done = wb.create_sheet(title="DONE")
        ws_done.append(["Idea Title", "Idea Brainstorm", "Link to Video", None, "YouTube Rating (/10)", "Calls Booked"])

        wb.save(path)

    def test_fetch_reads_ideas_from_sheet(self):
        connector = IdeaBankConnector(file_path=self.xlsx_path)
        result = connector.fetch()

        self.assertIsInstance(result, ConnectorResult)
        self.assertEqual(result.source, "idea_bank")
        self.assertEqual(len(result.items), 2)

        first = result.items[0]
        self.assertEqual(first.title, "Antigravity vs Claude Code Latency")
        self.assertIn("developer cognitive flow", first.body)
        self.assertEqual(first.source, "idea_bank")
        self.assertIn("Test_Idea_Bank.xlsx", first.url)

        second = result.items[1]
        self.assertEqual(second.title, "Kafka Consumer Lag War Story")
        self.assertIn("partition rebalancing", second.body)

    def test_fetch_ignores_empty_rows(self):
        # Add blank rows to the mock sheet
        wb = openpyxl.load_workbook(self.xlsx_path)
        ws = wb["IDEAS"]
        ws.append([None, None])
        ws.append(["", "   "])
        ws.append(["Valid Row After Gaps", "Valid brainstorm note"])
        wb.save(self.xlsx_path)

        connector = IdeaBankConnector(file_path=self.xlsx_path)
        result = connector.fetch()
        self.assertEqual(len(result.items), 3)
        self.assertEqual(result.items[-1].title, "Valid Row After Gaps")

    def test_append_idea_writes_to_sheet(self):
        connector = IdeaBankConnector(file_path=self.xlsx_path)
        row_idx = connector.append_idea(
            title="System Design: Modular Monolith vs Microservices",
            brainstorm="Why distributed boundaries should match data lifecycle, not team charts.",
        )
        self.assertTrue(row_idx >= 4)

        # Verify persisted to Excel
        wb = openpyxl.load_workbook(self.xlsx_path, data_only=True)
        ws = wb["IDEAS"]
        self.assertEqual(ws.cell(row=row_idx, column=1).value, "System Design: Modular Monolith vs Microservices")
        self.assertIn("team charts", ws.cell(row=row_idx, column=2).value)

    def test_append_packaging_writes_to_sheet(self):
        connector = IdeaBankConnector(file_path=self.xlsx_path)
        row_idx = connector.append_packaging(
            final_title="Stop Splitting Microservices Before You Hit 100k QPS",
            brainstorm="Modular monolith argument",
            template_inspiration="From [x] to [y] without [painful thing]",
            thumbnail_rec="Architecture Diagram Before/After",
        )
        self.assertTrue(row_idx >= 3)

        # Verify persisted to Excel
        wb = openpyxl.load_workbook(self.xlsx_path, data_only=True)
        ws = wb["PACKAGING"]
        self.assertEqual(ws.cell(row=row_idx, column=1).value, "Stop Splitting Microservices Before You Hit 100k QPS")
        self.assertEqual(ws.cell(row=row_idx, column=4).value, "From [x] to [y] without [painful thing]")

    def test_export_idea_batch(self):
        connector = IdeaBankConnector(file_path=self.xlsx_path)
        batch = [
            {"title": "Batch Idea 1", "brainstorm": "Note 1"},
            {"title": "Batch Idea 2", "brainstorm": "Note 2"},
            {"title": "Batch Idea 3", "brainstorm": "Note 3"},
        ]
        count = connector.export_idea_batch(batch)
        self.assertEqual(count, 3)

        # Verify fetch returns 2 original + 3 new = 5 items
        result = connector.fetch()
        self.assertEqual(len(result.items), 5)

    def test_export_packaging_batch(self):
        connector = IdeaBankConnector(file_path=self.xlsx_path)
        batch = [
            {
                "final_title": "Batch Pack 1",
                "brainstorm": "Note 1",
                "template_inspiration": "Template 1",
                "thumbnail_rec": "Thumb 1",
            },
            {
                "final_title": "Batch Pack 2",
                "brainstorm": "Note 2",
                "template_inspiration": "Template 2",
                "thumbnail_rec": "Thumb 2",
            },
        ]
        count = connector.export_packaging_batch(batch)
        self.assertEqual(count, 2)

        wb = openpyxl.load_workbook(self.xlsx_path, data_only=True)
        ws = wb["PACKAGING"]
        self.assertEqual(ws.cell(row=3, column=1).value, "Batch Pack 1")
        self.assertEqual(ws.cell(row=4, column=1).value, "Batch Pack 2")
        self.assertEqual(ws.cell(row=3, column=4).value, "Template 1")
        self.assertEqual(ws.cell(row=3, column=5).value, "Thumb 1")

    def test_missing_file_raises_file_not_found(self):
        connector = IdeaBankConnector(file_path=Path(self.temp_dir.name) / "nonexistent.xlsx")
        with self.assertRaises(FileNotFoundError):
            connector.fetch()


if __name__ == "__main__":
    unittest.main()
